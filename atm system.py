
from openpyxl import Workbook, load_workbook
import os

FILE = "atm_system_data.xlsx"

if not os.path.exists(FILE):
    wb = Workbook()
    ws = wb.active
    ws.append(["Name", "Mobile", "Balance"])
    wb.save(FILE)
    wb.close()

def create_account():
    wb = load_workbook(FILE)
    ws = wb.active

    name = input("Enter your name: ")
    mobile = input("Enter mobile number: ")

    for row in ws.iter_rows(min_row=2, values_only=True):
        if str(row[1]) == mobile:
            print("Account already exists!")
            wb.close()
            return

    ws.append([name, mobile, 500])
    wb.save(FILE)
    wb.close()

    print("Account created successfully!")
    print("Initial balance: Rs. 500")

def deposit():
    wb = load_workbook(FILE)
    ws = wb.active
    mobile = input("Enter mobile number: ")
    amount = float(input("Enter deposit amount: "))

    for row in ws.iter_rows(min_row=2):
        if str(row[1].value) == mobile:
            if amount > 0:
                row[2].value += amount
                wb.save(FILE)
                print("Deposit successful!")
                print("New balance:", row[2].value)
            else:
                print("Enter a valid amount!")
            wb.close()
            return

    wb.close()
    print("Account not found!")

def withdraw():
    wb = load_workbook(FILE)
    ws = wb.active
    mobile = input("Enter mobile number: ")
    amount = float(input("Enter withdrawal amount: "))

    for row in ws.iter_rows(min_row=2):
        if str(row[1].value) == mobile:
            if 0 < amount <= row[2].value:
                row[2].value -= amount
                wb.save(FILE)
                print("Withdrawal successful!")
                print("Remaining balance:", row[2].value)
            else:
                print("Invalid amount or insufficient balance!")
            wb.close()
            return

    wb.close()
    print("Account not found!")

def view_balance():
    wb = load_workbook(FILE)
    ws = wb.active
    mobile = input("Enter mobile number: ")

    for row in ws.iter_rows(min_row=2, values_only=True):
        if str(row[1]) == mobile:
            print("Name:", row[0])
            print("Mobile:", row[1])
            print("Balance: Rs.", row[2])
            wb.close()
            return

    wb.close()
    print("Account not found!")

while True:
    print("\n===== ATM SYSTEM =====")
    print("1. Create Account")
    print("2. Deposit Money")
    print("3. Withdraw Money")
    print("4. View Balance")
    print("5. Exit")

    choice = input("Enter your choice: ")

    if choice == "1":
        create_account()
    elif choice == "2":
        deposit()
    elif choice == "3":
        withdraw()
    elif choice == "4":
        view_balance()
    elif choice == "5":
        print("Thank you for using ATM!")
        break
    else:
        print("Invalid choice!")
