from openpyxl import Workbook, load_workbook
import os
FILE = "atm_data.xlsx"
if not os.path.exists(FILE):
    wb = Workbook()
    ws = wb.active
    ws.append(["Name", "Mobile", "Balance"])
    wb.save(FILE)
def create_account():
    wb = load_workbook(FILE)
    ws = wb.active
    name = input("Enter your name: ")
    mobile = input("Enter your mobile number: ")
    for row in ws.iter_rows(min_row=2, values_only=True):
        if str(row[1]) == mobile:
            print("Account already exists!")
            wb.close()
            return
    balance = 500
    ws.append([name, mobile, balance])
    wb.save(FILE)
    wb.close()
    print("Account created successfully!")
    print("Your initial balance is Rs.", balance)
def deposit():
    wb = load_workbook(FILE)
    ws = wb.active
    mobile = input("Enter your mobile number: ")
    amount = float(input("Enter deposit amount: "))
    if amount <= 0:
        print("Enter a valid amount!")
        wb.close()
        return
    for row in ws.iter_rows(min_row=2):
        if str(row[1].value) == mobile:
            row[2].value += amount
            new_balance = row[2].value
            wb.save(FILE)
            wb.close()
            print("Money deposited successfully!")
            print("New balance:", new_balance)
            return
    wb.close()
    print("Account not found!")
def withdraw():
    wb = load_workbook(FILE)
    ws = wb.active
    mobile = input("Enter your mobile number: ")
    amount = float(input("Enter withdrawal amount: "))
    if amount <= 0:
        print("Enter a valid amount!")
        wb.close()
        return
    for row in ws.iter_rows(min_row=2):
        if str(row[1].value) == mobile:
            if amount <= row[2].value:
                row[2].value -= amount
                new_balance = row[2].value
                wb.save(FILE)
                print("Withdrawal successful!")
                print("Remaining balance:", new_balance)
            else:
                print("Insufficient balance!")
            wb.close()
            return
    wb.close()
    print("Account not found!")
def view_balance():
    wb = load_workbook(FILE)
    ws = wb.active
    mobile = input("Enter your mobile number: ")
    for row in ws.iter_rows(min_row=2, values_only=True):
        if str(row[1]) == mobile:
            print("Name:", row[0])
            print("Mobile:", row[1])
            print("Balance: Rs.", row[2])
            wb.close()
            return
    wb.close()
    print("Account not found!")
while True:
    print("\n===== ATM SYSTEM =====")
    print("1. Create Account")
    print("2. Deposit Money")
    print("3. Withdraw Money")
    print("4. View Balance")
    print("5. Exit")
    choice = input("Enter your choice: ")
    if choice == "1":
        create_account()
    elif choice == "2":
        deposit()
    elif choice == "3":
        withdraw()
    elif choice == "4":
        view_balance()

    elif choice == "5":
        print("Thank you for using ATM!")
        break

    else:
        print("Invalid choice! Try again.")
